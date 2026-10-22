"""Builds Supplementary File S2: blank field-data workbook + data dictionary."""
import csv
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "field_protocol"

SHEETS = {
    "run_gates": [
        ("run_id", "string", "Unique run identifier (e.g. D03-A-S5-L50-R02)"),
        ("date_utc", "date", "Calendar date of the run (UTC)"),
        ("site_path", "string", "Approved site/path identifier (no operational IPs)"),
        ("authorization_ref", "string", "Reference to signed authorization (G1)"),
        ("isolation_checked", "bool", "Experimental segment and default-deny egress verified (G3)"),
        ("rollback_checked", "bool", "Independent rollback path verified (G3)"),
        ("clock_offset_ms", "float", "Measured clock offset before run (G6)"),
        ("operator_go", "bool", "Operator go decision"),
    ],
    "runs": [
        ("run_id", "string", "Run identifier"),
        ("block_id", "string", "Day/site/path randomization block"),
        ("scenario_id", "enum", "S0..S7"),
        ("load_level_pct", "enum", "20, 50 or 80 (% of measured stable throughput)"),
        ("measured_stable_throughput_mbps", "float", "Separately measured stable throughput of the path"),
        ("order_in_block", "int", "Randomized position within block"),
        ("t_baseline_start_utc", "datetime", "Start of 10-min baseline"),
        ("t_injection_start_utc", "datetime", "Start of perturbation (≤ 5 min)"),
        ("t_injection_end_utc", "datetime", "End of perturbation"),
        ("t_recovery_end_utc", "datetime", "End of 10-min recovery"),
        ("artifact_version", "string", "MetaTwin-Guard version / commit"),
        ("config_hash", "string", "SHA-256 of frozen configuration"),
        ("status", "enum", "valid | invalid | safety_stopped"),
    ],
    "events": [
        ("run_id", "string", "Run identifier"),
        ("event_id", "string", "Envelope event_id"),
        ("sensor_id", "string", "Pseudonymous sensor identity"),
        ("topology_node", "string", "Twin topology node"),
        ("source_seq", "int", "Source-local sequence"),
        ("event_time_utc", "datetime", "Event time"),
        ("ingest_time_utc", "datetime", "Ingest time"),
        ("twin_commit_time_utc", "datetime", "Twin commit time (Eq. 9)"),
        ("admitted", "bool", "Integrity admission g(e)"),
        ("reasons", "string", "Semicolon-separated rejection reasons"),
        ("robust_score", "float", "A(x), Eq. (5); empty if not evaluated"),
        ("rules_fired", "string", "Semicolon-separated static rules"),
        ("hybrid_alert", "bool", "Eq. (7)"),
        ("top3", "string", "Top-3 feature contributions"),
        ("ground_truth_injected", "bool", "Label from the injection schedule"),
        ("block_index", "int", "Evidence-chain block index"),
    ],
    "run_endpoints": [
        ("run_id", "string", "Run identifier"),
        ("mode", "enum", "rules | robust | hybrid | integrity_only | rules_integrity"),
        ("tp", "int", "True positives"), ("fp", "int", "False positives"),
        ("fn", "int", "False negatives"), ("tn", "int", "True negatives"),
        ("auprc", "float", "Area under precision-recall curve (continuous score only)"),
        ("false_alerts_per_hour", "float", "False alerts / normal exposure hours"),
        ("detection_latency_s", "float", "Eq. (8); empty if censored"),
        ("detection_censored", "bool", "True if no correct alert in window"),
        ("sync_latency_median_ms", "float", "Median Eq. (9) over admitted events"),
        ("recovery_time_s", "float", "Network recovery time"),
        ("evidence_chain_valid", "bool", "Ledger verification result"),
        ("cpu_pct_p95", "float", "95th percentile CPU utilization"),
        ("memory_mb_max", "float", "Peak resident memory"),
        ("storage_growth_mb", "float", "Evidence/log storage growth"),
    ],
    "accountability": [
        ("run_id", "string", "Run identifier"),
        ("status", "enum", "valid | invalid | safety_stopped"),
        ("stop_reason", "string", "Stop criterion triggered, if any"),
        ("replacement_run_id", "string", "Run scheduled as replacement (original stays listed)"),
        ("notes", "string", "Free text; no personal data"),
    ],
}


def main() -> None:
    OUT.mkdir(exist_ok=True)
    with (OUT / "S2_data_dictionary.csv").open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["sheet", "column", "type", "description"])
        for sheet, cols in SHEETS.items():
            for c in cols:
                w.writerow([sheet, *c])
    for sheet, cols in SHEETS.items():
        with (OUT / f"S2_{sheet}.csv").open("w", newline="") as fh:
            csv.writer(fh).writerow([c[0] for c in cols])
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("openpyxl not installed; CSV templates written only", file=sys.stderr)
        return
    wb = Workbook()
    wb.remove(wb.active)
    for sheet, cols in SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append([c[0] for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
    dd = wb.create_sheet("data_dictionary")
    dd.append(["sheet", "column", "type", "description"])
    for sheet, cols in SHEETS.items():
        for c in cols:
            dd.append([sheet, *c])
    wb.save(OUT / "S2_field_workbook.xlsx")
    print("S2 written to", OUT)


if __name__ == "__main__":
    main()
