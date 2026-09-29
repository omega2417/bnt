#!/usr/bin/env python3
"""Experiment dossier (English): archive contents, test-bench description, protocol,
visualised results and an Elsevier manuscript-integration kit.
-> report/DT_HIL_Experiment_Dossier_Testbed_EN.docx   (all numbers from results/*.json)"""
from __future__ import annotations

import csv
import json
import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
RES = ROOT / "results"
FIG = RES / "figures"
S = json.load(open(RES / "summary.json"))
MAN = json.load(open(RES / "manifest.json"))
HW = json.load(open(ROOT / "configs" / "hardware_bench_12v.json"))
PCFG = json.load(open(ROOT / "configs" / "policies.json"))
E5, PP, E11, T = S["e05"], S["e05"]["per_policy"], S["e11"], S["timing"]
NR = [r for r in S["runs"] if r["scenario_id"] not in ("E06", "E07")]
MAX_MAE = max(r["mae_soc_pp"] for r in NR)
MAX_EN = max(r["energy_error_pct"] for r in NR if r["energy_error_pct"] == r["energy_error_pct"])


def f(x, n=2):
    return f"{x:.{n}f}".replace("-", "−")


# ------------------------------------------------------------------ document setup
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.12
for lvl, size in ((1, 13), (2, 11.5), (3, 10.5)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Times New Roman"
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.italic = lvl == 3
    h.font.color.rgb = RGBColor(0, 0, 0)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True

# footer page numbers
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    r = fp.add_run()
    if tag:
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), tag)
        r._r.append(e)
    else:
        e = OxmlElement("w:instrText")
        e.set(qn("xml:space"), "preserve")
        e.text = text
        r._r.append(e)
    r.font.size = Pt(9)

FIGN, TABN = [0], [0]


def _runs(p, text, bold=False, italic=False, size=None, color=None, mono=False):
    for i, part in enumerate(text.split("**")):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if mono:
            r.font.name = "Consolas"
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Consolas")


def P(text="", bold=False, italic=False, size=None, align=None, color=None, after=None, mono=False):
    p = doc.add_paragraph()
    _runs(p, text, bold, italic, size, color, mono)
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def H(text, lvl=1):
    doc.add_heading(text, level=lvl)


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        _runs(p, it)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def code(lines):
    for ln in lines:
        p = P(ln, size=8.5, mono=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
        p.paragraph_format.left_indent = Cm(0.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def table(caption, header, rows, widths=None, note=None, size=8.5, mono_cols=()):
    TABN[0] += 1
    c = P(f"**Table {TABN[0]}.** {caption}", size=9.5, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
    c.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        shade(cell, "E8E7E3")
    repeat_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            _runs(cells[i].paragraphs[0], str(v), size=size, mono=i in mono_cols)
    if widths:
        for row in t.rows:
            for i, wdt in enumerate(widths):
                row.cells[i].width = Cm(wdt)
    if note:
        n = P(note, italic=True, size=8.5)
        n.paragraph_format.space_before = Pt(2)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return TABN[0]


def figure(name, caption, width=16.5):
    FIGN[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIG / f"{name}.png"), width=Cm(width))
    P(f"**Fig. {FIGN[0]}.** {caption}", size=9, after=8)
    return FIGN[0]


def box(lines, fill="F4F3F0"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    shade(cell, fill)
    cell.text = ""
    for k, ln in enumerate(lines):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        _runs(p, ln, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def human(n):
    for u in ("B", "KB", "MB"):
        if n < 1024:
            return f"{n:.0f} {u}" if u == "B" else f"{n:.1f} {u}"
        n /= 1024
    return f"{n:.1f} GB"


def folder_stats(rel):
    p = ROOT / rel
    files = [q for q in p.rglob("*") if q.is_file() and "__pycache__" not in q.parts and "_UA" not in q.name and "_ua" not in q.name]
    return len(files), sum(q.stat().st_size for q in files)


# ============================================================ TITLE
P("Supporting document for an Elsevier manuscript · Applied Energy", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.LEFT, color="52514E")
P("Experiment Dossier: Software–Hardware Validation of the Digital Twin of a Communication-Node Power Supply",
  bold=True, size=16, align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
P("Archive contents · test-bench description · experimental protocol · visualised results and tables · manuscript integration kit",
  size=11, align=WD_ALIGN_PARAGRAPH.LEFT, color="333333", after=6)
P(f"Package: dt_power_hil (dtpower {S['dtpower_version']}) · {MAN['n_runs']} runs · configuration SHA-256 {S['config_sha256'][:16]}… · "
  f"generated 29 September 2026", size=9, align=WD_ALIGN_PARAGRAPH.LEFT, color="52514E")
P("[Authors and affiliations — to be completed from the manuscript]", italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT)
box(["**EVIDENCE CLASSIFICATION.** Software-in-the-loop (SIL) simulation with emulated hardware (virtual testbed). The battery, DC bus, "
     "load branches, INA226/INA228 measurement channels, MQTT link and ESP32 controller are reproduced by models parameterised with "
     "datasheet characteristics of the real components listed in Part B. **No physical measurement was performed.** The results must be "
     "reported as a pre-hardware (virtual-testbed) validation, never as HIL, CHIL, PHIL or field results. Part B describes the physical bench "
     "that the same software and protocols are prepared for."], fill="FBF0D9")

H("How to use this document")
table("Structure of the dossier and its intended use in the manuscript",
      ["Part", "Content", "Use in the Elsevier manuscript"],
      [["A. Archive contents", "Directory inventory, file catalogue, data dictionary, run registry, reproduction and integrity",
        "Data availability statement; supplementary material description"],
       ["B. Test bench", "Architecture, wiring, components with characteristics, measurement points, measurement chain, controller, timing, protection, calibration, emulation parameters",
        "New subsection “Experimental testbed” in Methods (text and Tables B-x can be pasted)"],
       ["C. Protocol", "Hypotheses, scenarios, policies, metric definitions, statistical design", "Methods: “Experimental protocol and metrics”"],
       ["D. Results", "Acceptance summary, metrology, scenario dynamics, twin accuracy, timing, faults, policy comparison, 14-day reproduction, sizing",
        "Results: “Validation on the virtual testbed”; figures and tables with captions"],
       ["E. Integration kit", "Ready-to-paste section text, figure/table mapping, abstract/highlight/conclusion insertions, data statement, reviewer response, allowed claims",
        "Direct insertion and revision letter"],
       ["Appendix", "Full per-run metrics (51 runs), calibration levels", "Supplementary tables"]],
      widths=[3.2, 7.6, 6.2])

# ============================================================ PART A
page_break()
H("Part A. Archive contents")
H("A.1. Package overview", 2)
n_all, s_all = folder_stats(".")
P(f"The archive `dt_power_hil_EN.zip` unpacks into a single directory `dt_power_hil/` ({n_all} files, {human(s_all)} uncompressed). "
  "It is self-contained: one command regenerates every result, figure, protocol, report and the web page from the configuration files. "
  "Fig. 1 shows the data flow from inputs to deliverables.")
figure("fig13_data_pipeline", "Data pipeline of the package: configuration → campaign runner (dtpower) → result package (summary, raw logs, "
       "protocols, calibration) → generators of figures, measurement protocols, reports and the web dashboard. Tests and firmware mirror the controller logic.")

H("A.2. Directory inventory", 2)
inv = [("dtpower/", "Python package: digital twin, virtual testbed, EMS, statistics", "hand-written source"),
       ("scripts/", "Campaign runner and generators of figures, protocols, reports, web", "hand-written source"),
       ("tests/", "Software acceptance tests (protocol §18.3)", "hand-written source"),
       ("configs/", "Hardware passport, policy thresholds, manuscript scenarios S1–S3", "hand-written input"),
       ("firmware/", "ESP32-S3 edge_guard firmware (ESP-IDF, C)", "hand-written source, not compiled"),
       ("hardware/", "Bill of materials, wiring, measurement protocols P-1…P-7", "BOM hand-written; protocols generated"),
       ("results/", "Complete result package of the emulated campaign", "generated by run_campaign.py"),
       ("results/raw/", "Per-run raw logs (telemetry, events, commands, edge log)", "generated"),
       ("results/protocols/", "Per-run protocols in the §21.3 format", "generated"),
       ("results/calibration/", "E00 calibration, E09 capacity, channel characteristics, uncertainty budget", "generated"),
       ("results/figures/", "Figures fig01–fig13 as PNG (300 dpi) and vector PDF", "generated"),
       ("results/configs/", "Frozen copy of the configurations used by the campaign", "generated"),
       ("report/", "Elsevier-style report and this dossier (DOCX)", "generated"),
       ("web/", "Self-contained results dashboard (index.html) and its template", "generated / template")]
rows = []
for rel, what, how in inv:
    n, s = folder_stats(rel)
    rows.append([rel, what, n, human(s), how])
table("Directory inventory of the archive (Ukrainian-language files excluded)", ["Directory", "Content", "Files", "Size", "Origin"], rows,
      widths=[3.0, 7.0, 1.3, 1.7, 4.0], mono_cols=(0,))

H("A.3. File catalogue", 2)
cat = [
    ("README.md", "Package description, reproduction commands, migration steps to the physical bench"),
    ("configs/hardware_bench_12v.json", "Hardware passport: battery, BMS limits, generation source, branches, INA226/INA228/DS18B20 characteristics, shunts, channels M0–M5/REF, calibration standard, timing budget, network model, protection"),
    ("configs/policies.json", "Frozen thresholds of policies C0, C1, C2, C_oracle; stop threshold SOC_ref = 0.20; ε_P = 0.5 W"),
    ("configs/article_scenarios.json", "Manuscript parameters of S1–S3 (Tables 2–3), 14-day PSH array, values reported in the manuscript"),
    ("dtpower/__init__.py", "Package version and evidence-class string"),
    ("dtpower/config.py", "Configuration loading, SHA-256 hashing, run-configuration validation (§21.4 rejection rules)"),
    ("dtpower/profiles.py", "Normalised PV shape, lighting mask, hourly profiles, zero-order-hold upsampling, step profiles"),
    ("dtpower/sizing.py", "Analytical sizing (B.4–B.7), sensitivity table and grid"),
    ("dtpower/energy_sim.py", "Energy-level SOC simulator of the manuscript with pre-step limits, unserved and curtailed energy bookkeeping"),
    ("dtpower/battery.py", "LiFePO₄ equivalent-circuit plant (OCV table, R0, R1C1, η_Q, BMS cut-offs) with hidden unit parameters"),
    ("dtpower/sensors.py", "Measurement-chain emulation: shunt tolerance/TCR, IC offset, gain error, noise, LSB quantisation, saturation; DS18B20"),
    ("dtpower/calibration.py", "E00 procedure (zero, 5 levels, fit/verify split), uncertainty budget, protocol-row derivation"),
    ("dtpower/capacity.py", "E09 working-interval capacity and coulombic-efficiency test"),
    ("dtpower/transport.py", "MQTT-like channel: log-normal latency, loss, QoS-1 duplicates, blackouts"),
    ("dtpower/edge_guard.py", "Python mirror of the ESP32 firmware: state machine, command contract, watchdog, local protection"),
    ("dtpower/twin.py", "Telemetry validator; SOC prior/posterior (coulomb counting + EKF); bus balance; open-loop prediction"),
    ("dtpower/forecast.py", "Causal generation and demand forecasters (no access to the scenario profile)"),
    ("dtpower/ems.py", "Policies C0 (schedule), C1 (SOC hysteresis), C2 (forecast rule) and C_oracle"),
    ("dtpower/bench.py", "Virtual-testbed orchestrator (plant, sensors, edge, transport, twin, EMS, REF logger, service probe) and bench builder"),
    ("dtpower/scenarios.py", "Scenario matrix E01–E08"),
    ("dtpower/replay.py", "E11: SIL replication of S1–S3, reconciliation grid, lockstep controller replay"),
    ("dtpower/stats.py", "BCa bootstrap, exact sign-flip test, Holm correction, Friedman test"),
    ("scripts/run_campaign.py", "Runs E00–E11 in parallel, writes raw logs, protocols, metrics, summary, manifest, checksums"),
    ("scripts/make_figures.py", "Figures fig01–fig09"),
    ("scripts/make_figures_extra.py", "Figures fig10–fig13 (wiring, scenario dynamics, uncertainty budget, pipeline)"),
    ("scripts/make_protocols.py", "Generates hardware/measurement_protocols.md"),
    ("scripts/make_report.py", "Generates the Elsevier-style report"),
    ("scripts/make_experiment_book.py", "Generates this dossier"),
    ("scripts/make_web.py", "Builds web/index.html with embedded data"),
    ("tests/test_acceptance.py", "14 acceptance checks: 10 W·1 h = 10 Wh, sign convention, no energy creation, curtailment, night PV = 0, manuscript sizing, stale/duplicate/boot-id commands, no-future forecaster, validator, config rules, statistics, end-to-end run"),
    ("firmware/esp32_edge_guard/main/edge_guard.c", "ESP-IDF firmware: INA226 over I²C, 10 Hz sampling, 1 Hz telemetry, state machine, command contract, watchdog, relay GPIO"),
    ("firmware/esp32_edge_guard/CMakeLists.txt, main/CMakeLists.txt, main/Kconfig.projbuild", "ESP-IDF build files; broker URI option"),
    ("hardware/BOM_and_wiring.md", "24-item bill of materials, functional diagram, meter selection check, installation rules"),
    ("hardware/measurement_protocols.md", "Protocols P-1…P-7 with EMU values and blank PHYS columns"),
    ("results/summary.json", "All metrics, statistics, calibration, E06/E07/E10/E11 outputs, acceptance table"),
    ("results/metrics_runs.csv", "One row per run: 40 metric columns (see Table A-4)"),
    ("results/series_selected.json", "10-s time series of selected runs used by figures and the web page"),
    ("results/manifest.json", "Versions, seeds, run list, wall time, evidence class"),
    ("results/checksums.sha256", "SHA-256 of every file in results/"),
    ("report/DT_HIL_Experiment_Report_Elsevier_EN.docx", "Elsevier-style experiment report"),
    ("report/DT_HIL_Experiment_Dossier_Testbed_EN.docx", "This dossier"),
    ("web/index.html", "Self-contained interactive dashboard (opens offline in a browser)"),
]
table("File catalogue (excluding per-run files, which follow the pattern in Table A-3)", ["Path", "Description"], cat,
      widths=[6.0, 11.0], size=8, mono_cols=(0,))

H("A.4. Data dictionary", 2)
table("Per-run raw files in results/raw/<run_id>/",
      ["File", "Rows", "Columns and meaning"],
      [["telemetry_10s.csv", "every 10 s + every relay switch",
        "t (s, model time); soc_true (hidden plant SOC, emulation only); soc_ref (independent REF SOC); soc_prior / soc_post (twin before/after update); "
        "v_batt (V); i_batt (A, + = discharge); i_ref (A, REF channel); p_gen_avail (W, scenario); p_gen (W, accepted by the bus); p_crit_term, p_aux_term (W at load terminals); "
        "aux_sched (0/1 schedule); aux_cmd (0/1 ESP32 output); relay (0/1 physical state); pred_min (C2 forecast min SOC); valid (cumulative valid samples); "
        "residual_w (bus balance r_P, W); service_ok (0/1 HTTP probe)"],
       ["events.csv", "per event", "run_id, event_id, t_model_s, event_type (RUN_START, TELEMETRY_STALE, TELEMETRY_RESTORED, INVALID_SAMPLE, STOP_LIMIT_SOC_REF, RUN_END), value"],
       ["commands.csv", "per command", "run_id, command_id, action (SET_AUX_ON/OFF), reason, sent_s, ack_s, effect_s, effect_kind (edge / state_confirmed), accepted, reject_reason, retry_of"],
       ["edge_log.csv", "per controller event", "t_s, kind (AUX_ON, AUX_OFF, ACK, REJECT, STATE), detail"],
       ["results/protocols/<run_id>.txt", "1 per run", "Run protocol fields of §21.3 (RUN_ID … ANALYST_REVIEW)"]],
      widths=[3.4, 2.6, 11.0], size=8, mono_cols=(0,))
with open(RES / "metrics_runs.csv") as fh:
    cols = next(csv.reader(fh))
desc = {"run_id": "unique run identifier <date>_<scenario>_<policy>_R<repeat>_EMU", "scenario_id": "E01…E08", "policy": "C0, C1, C2, C_oracle",
        "repeat": "block index (E05: 0–9)", "mode": "sil_hw_emulated", "duration_h": "run duration", "stop_reason": "completed / stopped_limit",
        "right_censored": "true if the limit was not reached", "coverage_pct": "valid telemetry share", "unknown_duration_s": "time without telemetry beyond 2 s",
        "soc_ref_start": "initial SOC_ref", "min_soc_ref": "minimum SOC_ref", "mae_soc_pp": "MAE posterior vs SOC_ref, pp", "rmse_soc_pp": "RMSE, pp",
        "max_soc_err_pp": "max |error|, pp", "ref_vs_true_mae_pp": "REF vs hidden truth (emulation only)", "energy_error_pct": "twin vs REF discharge energy, %",
        "critical_unserved_wh": "critical unserved energy", "aux_served_wh": "lighting energy delivered", "aux_unserved_wh": "lighting unserved (incl. EMS curtailment)",
        "lolh_critical_h": "loss-of-load hours, critical", "lpsp_critical": "LPSP_E critical", "lpsp_total": "LPSP_E total", "service_checks": "HTTP probes",
        "service_failures": "failed probes", "availability_pct": "observed availability", "commands": "commands sent", "commands_accepted": "accepted by ESP32",
        "commands_effect_confirmed": "accepted with confirmed effect", "switch_count": "relay transitions", "residual_w_mean": "mean bus residual r_P",
        "residual_w_sd": "SD of r_P", "e_gen_wh": "generation accepted", "e_curtailed_wh": "generation curtailed", "e_ref_dis_wh": "discharge energy by REF",
        "rtt_median_ms": "command RTT median", "rtt_p95_ms": "RTT p95", "rtt_max_ms": "RTT max", "effect_median_ms": "decision → current edge median",
        "effect_max_ms": "decision → current edge max"}
table("Columns of results/metrics_runs.csv", ["Column", "Meaning"], [[c, desc.get(c, "")] for c in cols], widths=[4.2, 12.8], size=8, mono_cols=(0,))
table("Top-level keys of results/summary.json",
      ["Key", "Content"],
      [["evidence_class, dtpower_version, config_sha256", "identification"], ["bench", "seed and hidden truth of the emulated unit"],
       ["calibration, capacity, channels, uncertainty", "E00, E09, P-1 rows, P-3 budget"], ["runs", "per-run metrics (as metrics_runs.csv, plus transport statistics)"],
       ["e05", "per-policy summaries, primary/secondary paired contrasts, Friedman, common-interval analysis"], ["e06, e07, e10", "fault-handling results"],
       ["e11", "SIL S1–S3, 1-min check, reconciliation grid, controller replay"], ["sizing, sensitivity, sensitivity_grid", "analytical sizing"],
       ["timing", "RTT and decision-to-effect statistics"], ["acceptance", "protocol §17 criteria with verdicts"]],
      widths=[5.5, 11.5], size=8, mono_cols=(0,))

H("A.5. Run registry", 2)
reg = {}
for r in S["runs"]:
    reg.setdefault(r["scenario_id"], []).append(r)
table("Registry of the 51 runs by scenario",
      ["Scenario", "Runs", "Policies", "Seeds", "Duration, h", "Stop reasons"],
      [[sid, len(rs), ", ".join(sorted({r["policy"] for r in rs})), f"{min(r['seed'] for r in rs)}…{max(r['seed'] for r in rs)}",
        f"{f(min(r['duration_h'] for r in rs))}–{f(max(r['duration_h'] for r in rs))}",
        ", ".join(f"{k}×{v}" for k, v in sorted({x: sum(1 for r in rs if r['stop_reason'] == x) for x in {r['stop_reason'] for r in rs}}.items()))]
       for sid, rs in sorted(reg.items())],
      widths=[2.0, 1.4, 3.8, 3.8, 2.4, 3.6], size=8.5)

H("A.6. Reproduction and integrity", 2)
code(["pip install numpy scipy matplotlib python-docx pytest",
      "cd dt_power_hil",
      "python -m pytest -q tests                  # 14 acceptance checks",
      "python scripts/run_campaign.py             # E00–E11 → results/  (≈ 80 s, 4 cores)",
      "python scripts/make_figures.py && python scripts/make_figures_extra.py",
      "python scripts/make_protocols.py && python scripts/make_report.py",
      "python scripts/make_experiment_book.py && python scripts/make_web.py",
      "sha256sum -c results/checksums.sha256      # integrity check (run inside results/)"])
P(f"Environment of the reference run: Python {MAN['python']}, NumPy {MAN['numpy']}, platform {MAN['platform']}. Seeds: bench {MAN['bench_seed']}, "
  f"base {MAN['base_seed']}; E05 blocks use seed base + 1000·(block+1), identical for all policies of a block (paired design).")

# ============================================================ PART B
page_break()
H("Part B. Test-bench description")
P("This part is written so that it can be inserted into the Methods section of the manuscript as the subsection “Experimental testbed”. "
  "It describes the physical 12 V bench on real components; the same description parameterises the emulation used in Part D.", italic=True)
H("B.1. Purpose and scope", 2)
P("The bench closes the loop between the digital twin and a battery-supplied communication node: real (or emulated) measurements feed the twin, "
  "the energy-management system (EMS) decides whether the auxiliary lighting may run, and the decision is executed by a relay and confirmed by a "
  "measured change of current. The bench implements the mandatory programme of the protocol: one 12 V class stand, measured generation, "
  "battery and load channels, an independent reference, closed-loop control of the auxiliary load, repeatable disturbances, service availability "
  "monitoring and immutable logs. Its 12.8 V / 40 Ah battery (≈ 512 Wh) validates the loop and the model of this bench; it is not a physical "
  "replica of manuscript scenario S1 (12 V / 200 Ah), and scale transfer requires a separate scaling model.")
H("B.2. Architecture and wiring", 2)
figure("fig01_architecture", "Functional architecture: hardware part (left) and software part (right). The EMS never receives the future scenario profile.")
figure("fig10_bench_wiring", "Wiring of the bench with fuses, shunts, measurement points M0–M5 and REF, the switched auxiliary branch, "
       "the never-switched critical branch, the ESP32 edge controller, the PC and the independent Raspberry Pi logger.", 17.0)
H("B.3. Components and characteristics", 2)
bom = [["Controller", "Espressif ESP32-S3-DevKitC-1 (ESP32-S3-WROOM-1-N8R8)", "2×240 MHz, 3.3 V logic, esp_timer 1 µs; 5 V USB from PC UPS", "1"],
       ["Battery", "LiFePO₄ 4S1P 12.8 V / 40 Ah, DALY 4S 12 V 40 A-class BMS", "CV 14.2 V; OVP 14.6 V; UVP 10.0 V; I_ch ≤ 20 A; I_dis ≤ 40 A; T_ch 0–45 °C", "1"],
       ["Generation G1", "RIDEN RD6018 + Mean Well RSP-750-48", "0–60 V, 0–18 A, Modbus RTU; 0.01 V / 0.01 A set resolution", "1+1"],
       ["Reverse-current block", "LTC4359 ideal-diode module", "minimal forward loss at M0", "1"],
       ["Critical DC/DC", "Mean Well DDR-60G-12", "in 9–36 V, out 12 V / 5 A, η ≈ 0.90", "1"],
       ["Communication node", "MikroTik hAP ac² (local HTTP service)", "consumption measured by M3", "1"],
       ["Load top-up", "Korad KEL103 (or Rigol DL3021), CP mode", "node + e-load = 18 W at terminals", "1"],
       ["Auxiliary DC/DC", "Mean Well DDR-60G-12", "as above", "1"],
       ["Auxiliary load", "12 V / 30 W LED floodlight", "measured by M5", "1"],
       ["Auxiliary switch", "Bosch 0 332 019 150 relay (NO) + PC817 + ULN2003A + 1N4007", "operate ≈ 10 ms; open (OFF) when unpowered", "1"],
       ["Meters M0–M5", "TI INA226", "16-bit; shunt LSB 2.5 µV (±81.92 mV); bus LSB 1.25 mV (≤ 36 V); offset ≤ ±10 µV; gain ≤ 0.1 %", "6"],
       ["Reference REF", "TI INA228 on Raspberry Pi 4", "20-bit; shunt LSB 312.5 nV (±163.84 mV); bus LSB 195.3 µV (≤ 85 V); offset ≤ ±1 µV; gain ≤ 0.05 %", "1"],
       ["Shunts", "FL-2 50 A / 75 mV (M1, REF); 30 A / 75 mV (M0); 10 A / 75 mV (M2–M5), class 0.5", "1.5 / 2.5 / 7.5 mΩ", "7"],
       ["Temperature", "Maxim DS18B20", "±0.5 °C (−10…+85 °C), 0.0625 °C resolution", "2"],
       ["I²C isolation", "Analog Devices ADuM1250", "breaks USB ↔ bus ground loop", "1"],
       ["Service probe N1", "Raspberry Pi 4 Model B on mains", "HTTP request every 1 s", "shared"],
       ["Protection", "MIDI 40 A; ATO 25 / 5 / 7.5 A; Blue Sea Systems 6006", "DC rated fuses and battery disconnect", "—"],
       ["Calibration standard", "Fluke 87V + certified 4-terminal shunt", "expanded uncertainty from certificate (placeholder 0.05 %)", "1"]]
table("Bench components and characteristics used in the analysis (nominal datasheet values; verify against the actual units)",
      ["Function", "Component", "Characteristics", "Qty"], bom, widths=[2.8, 5.6, 7.6, 1.0], size=8)
H("B.4. Measurement points", 2)
LOC = {"M0": "bus after the generation source", "M1": "battery branch, signed current", "M2": "critical branch input",
       "M3": "critical load terminals", "M4": "auxiliary branch input", "M5": "auxiliary load terminals", "REF": "independent battery channel"}
PURP = {"M0": "power injected into the bus", "M1": "twin state estimation, charge/discharge energy", "M2": "bus-side critical demand",
        "M3": "useful critical energy", "M4": "lighting power and physical switching confirmation", "M5": "useful lighting energy",
        "REF": "reference SOC and check of M1"}
table("Measurement points and their role",
      ["Point", "Location", "Quantities", "Purpose"],
      [[k, LOC[k], "V, I, P" if k != "REF" else "I, V, ∫I dt", PURP[k]] for k in LOC] +
      [["T1/T2", "battery case / ambient air", "temperature", "test conditions, charge-temperature limits"],
       ["N1", "independent network client", "request success, RTT", "service availability of the node"]],
      widths=[1.4, 5.0, 2.6, 8.0], size=8.5)
H("B.5. Measurement chain and selection check", 2)
table("Measurement-channel characteristics derived from datasheets (protocol P-1)",
      ["Channel", "Meter / shunt", "Full scale, A", "LSB, mA", "Max offset, mA", "Uncal. error, %", "Max shunt P, W", "Calibration levels, A"],
      [[r["channel"], f"{r['meter']} / {f(r['r_shunt_mohm'], 1)} mΩ", f(r["i_full_scale_a"], 1), f(r["i_lsb_ma"], 3), "±" + f(r["offset_max_ma"], 2),
        "±" + f(r["uncal_gain_err_max_pct"], 2), f(r["p_shunt_max_w"], 2), "; ".join(f(x, 2) for x in r["calibration_levels_a"])] for r in S["channels"]],
      widths=[1.3, 3.3, 1.7, 1.5, 1.9, 1.8, 1.8, 3.7], size=7.8,
      note="The INA226 common-mode limit (36 V) suits the 12.8 V bus but not S3 (51.2 V); an uncorrected ±10 µV offset on 1.5 mΩ is ±6.7 mA, i.e. up to 0.4 pp SOC drift per 24 h, "
           "so zero calibration is mandatory. Sign convention: positive battery current = discharge.")
H("B.6. Controller, communication and timing budget", 2)
tm = HW["timing"]
table("Timing budget and communication parameters",
      ["Item", "Value", "Remark"],
      [["Local sampling (INA226 ×6)", f"{tm['sample_hz']} Hz", "averaging 16, 1.1 ms conversion; energy changes, not switching transients"],
       ["Aggregated telemetry", f"{tm['telemetry_period_s']} s", "mean / min / max / n per channel; key (run_id, device_id, boot_id, seq)"],
       ["Twin update", f"{tm['twin_period_s']} s", "prior before, posterior after each telemetry"],
       ["EMS decision", f"{tm['ems_period_s']} s", "urgent OFF allowed earlier"],
       ["Heartbeat / stale threshold", f"{tm['heartbeat_s']} s / {tm['telemetry_stale_s']} s", "DEGRADED and aux OFF on expiry"],
       ["Command TTL / ACK timeout", f"{tm['command_ttl_ms']} ms / {tm['ack_timeout_s']} s", "idempotent SET_AUX_ON / SET_AUX_OFF; TOGGLE forbidden"],
       ["Clock uncertainty", f"±{tm['clock_uncertainty_ms']} ms", "SNTP; used conservatively in the expiry check"],
       ["MQTT topics", "dt/node01/{telemetry, state, command, ack, event, health}", "no retained commands; QoS 1 for telemetry and commands"],
       ["Controller states", "INIT → READY → RUNNING ⇄ DEGRADED → STOPPED", "aux OFF in every state except RUNNING"]],
      widths=[4.2, 5.0, 7.8], size=8.5)
H("B.7. Protection and safety", 2)
bullets(["Battery fuse MIDI 40 A within 150 mm of the “+” terminal; ATO 25 A (generation), 5 A (critical), 7.5 A (auxiliary) in DC-rated holders.",
         "Blue Sea 6006 disconnect opens the power circuit; PC, ESP32 and Raspberry Pi run from a UPS and keep logging.",
         "The auxiliary relay is normally open and GPIO10 is pulled down, so the auxiliary branch is OFF at boot, reset or loss of control.",
         "The critical branch has no controlled switch; it is protected only by its fuse and the BMS.",
         "BMS and hardware limits have priority over all software thresholds; no deliberate short circuits, BMS bypass, overcharge or deep discharge; boundary tests (E10) by signal injection only.",
         "Physical runs stop at SOC_ref = 0.20 (not the model SOC_min = 0.15)."])
H("B.8. Calibration and independent reference", 2)
cap = S["capacity"]
P(f"All channels are calibrated in E00 (60-s zero, five levels, fit on levels 1/3/5, verification on 2/4) against a certified standard. The "
  f"reference SOC_ref is computed by the Raspberry Pi logger from the REF channel with its own shunt, ADC, clock and coefficients, and with the "
  f"capacity Q_ref and coulombic efficiency η_Q measured in E09 (CC-CV 20 A / 14.2 V, 2 A tail; 10 A discharge to 11.2 V). In the emulated "
  f"campaign E09 gave Q_ref = {f(cap['q_ref_ah'], 3)} Ah, Q_M1 = {f(cap['q_m1_ah'], 3)} Ah, E_ref = {f(cap['e_ref_wh'], 1)} Wh and η_Q = {f(cap['eta_q_ref'], 4)}.")
H("B.9. Software stack", 2)
table("Software components of the bench",
      ["Layer", "Implementation", "Status"],
      [["Controller firmware", "ESP-IDF v5.x, C (edge_guard.c): I²C master, esp_timer 10 Hz, ESP-MQTT, cJSON", "written; to be compiled and E10-tested on target"],
       ["Broker", "Mosquitto 2.x on the PC, isolated bench VLAN", "configuration step"],
       ["Twin / EMS / logger", f"Python {MAN['python']} package dtpower {S['dtpower_version']}", "implemented and tested (14 tests)"],
       ["Analysis", f"NumPy {MAN['numpy']}, SciPy, Matplotlib; BCa, sign-flip, Holm, Friedman", "implemented"],
       ["Reference logger", "Raspberry Pi 4, INA228 driver, local CSV", "to be implemented with the physical bench"],
       ["Generation driver", "RD6018 Modbus RTU set-points every 1 s", "to be implemented with the physical bench"]],
      widths=[3.5, 8.5, 5.0], size=8.5)
H("B.10. Emulation of the real components (used for the campaign in Part D)", 2)
bt = HW["battery"]["emulation_truth"]
tr = S["bench"]["hidden_truth"]
table("How each real element was emulated and with which parameters",
      ["Real element", "Emulation model", "Parameters (drawn value for bench unit, seed 1)"],
      [["LiFePO₄ pack + BMS", "ECM: OCV(z)+shift, R0, R1C1, η_Q; BMS opens outside 10.0–14.6 V",
        f"capacity N({bt['capacity_ah_mean']}, {bt['capacity_ah_sd']}) → {f(tr['capacity_ah'], 2)} Ah; R0 → {f(tr['r0'] * 1e3, 1)} mΩ; R1 = {bt['r1_ohm'] * 1e3:.0f} mΩ, C1 = {bt['c1_f']:.0f} F; η_Q = {bt['coulombic_eff']}; OCV shift {f(tr['ocv_shift_v'] * 1e3, 1)} mV"],
       ["RD6018 + LTC4359 (G1)", "power source with CV 14.2 V acceptance and 20 A charge limit; curtailment booked", "P_avail from scenario profile"],
       ["DDR-60G-12 branches", "constant efficiency per unit", f"η_crit = {f(S['bench']['eta_crit'], 4)}, η_aux = {f(S['bench']['eta_aux'], 4)}"],
       ["Node + e-load", "18 W ± 2 % random ripple", "per-second draw"],
       ["Bosch relay + LED", "switch delay 10 ms + 150 ms soft start; 10 Hz edge detection", "OFF default"],
       ["INA226 / INA228 + FL-2", "shunt class and gain/offset errors uniform within datasheet maxima; Gaussian noise; LSB rounding; saturation", "see Table B-5; corrected by E00 coefficients"],
       ["Bus self-consumption", "unmetered constant load", f"{HW['other_bus_load_w']} W"],
       ["MQTT / Wi-Fi", "log-normal one-way latency, loss, QoS-1 duplicates, blackouts", f"median {HW['network']['latency_median_ms']} ms, σ_ln {HW['network']['latency_sigma']}, loss {HW['network']['loss_prob']}, dup {HW['network']['dup_prob']}"],
       ["ESP32 edge_guard", "Python mirror of edge_guard.c", "clock offset uniform ±15 ms"],
       ["Network service", "probe succeeds if the DC/DC input ≥ 9 V and BMS closed; random non-power failures", f"p_fail = {HW['network']['service_check_fail_prob']}"]],
      widths=[3.4, 6.4, 7.2], size=8)

# ============================================================ PART C
page_break()
H("Part C. Experimental protocol")
table("Hypotheses and required evidence (protocol §1.1)",
      ["Code", "Hypothesis", "Evidence used"],
      [["H1", "The model reproduces the energy balance of the bench", "SOC error vs independent REF; discharge-energy error; bus residual"],
       ["H2", "The twin stays current under disturbances", "telemetry coverage, age, stale detection, RTT"],
       ["H3", "Twin decisions change the physical state", "chain decision → command → ACK → measured current edge"],
       ["H4", "Predictive control improves a pre-specified resilience metric", "paired E05 blocks; lighting delivered co-reported"]],
      widths=[1.3, 7.0, 8.7])
sc_rows = [["E00/E09", "Calibration of 7 channels; working-interval capacity", "—", "—"],
           ["E01", "Stabilisation 25 W; gen 10 W < demand; gen 60 W > demand; lighting off", "60 min", "C0"],
           ["E02", "18 W base; +30 W LED at 15 min, off at 30, on at 45; gen 25 W", "60 min", "C0"],
           ["E03", "Gen 60 → 10 → 60 W (30/60/30 min); lighting scheduled", "120 min", "C0, C1, C2"],
           ["E04", "Loss of 55 W source for 60 min", "90 min", "C0"],
           ["E05", "SOC₀ 0.45; gen 40 W 1.5 h / 5 W 3 h / 60 W 1.5 h; 18 + 30 W; stop SOC_ref 0.20", "≤ 6 h", f"C0, C1, C2, C_oracle × {E5['blocks']}"],
           ["E06", "Uplink loss 30 s; bidirectional loss 30 s; stale and foreign-run commands", "30 min", "C1"],
           ["E07", "Faults in twin telemetry copy: offset, stuck, missing, impossible, sign", "30 min", "C1"],
           ["E08", "SOC₀ 0.33; 20-min deficit; recovery 120 W", "60 min", "C0, C1, C2"],
           ["E10", "12 protection / contract vectors", "—", "—"],
           ["E11", "14 days S1–S3 (1 h / 1 min); lockstep replay S1 12.8 V / 100 Ah", "336 model h", "C0, C1, C2"]]
table("Scenario matrix (all executed in mode sil_hw_emulated)", ["Code", "Scenario", "Duration", "Policies"], sc_rows, widths=[1.6, 10.0, 2.2, 3.2], size=8.5)
c1, c2 = PCFG["C1"], PCFG["C2"]
table("Control policies (thresholds frozen before the campaign)",
      ["Policy", "Rule", "Parameters"],
      [["C0", "lighting by schedule; shared protections only", "—"],
       ["C1", "SOC hysteresis", f"off at SOC ≤ {c1['aux_off_soc']}, on at SOC ≥ {c1['aux_on_soc']}, dwell {c1['minimum_dwell_s']} s"],
       ["C2", "forecast rule: simulate the horizon with lighting; switch off if forecast min SOC < reserve",
        f"horizon {c2['horizon_s'] // 3600} h, step {c2['step_s'] // 60} min, reserve {c2['reserve_soc']}, return at forecast ≥ {c2['return_pred_soc']} and SOC ≥ {c2['return_now_soc']}; causal exponential-smoothing forecast"],
       ["C_oracle", "C2 logic with the true future generation (information bound only)", "never mixed with C2"]],
      widths=[1.6, 6.4, 9.0], size=8.5)
table("Metric definitions (protocol §16)",
      ["Metric", "Definition"],
      [["MAE / RMSE SOC", "100·mean|ẑ − z_ref| and 100·√mean(ẑ − z_ref)², percentage points; ẑ = twin posterior, z_ref from REF"],
       ["Energy error ε_E", "100·|E_twin − E_ref| / E_ref on the discharge energy"],
       ["Bus residual r_P", "P_gen + V·I_batt − P_crit − P_aux (W), measured channels"],
       ["Unserved energy", "∫max(0, P_dem − P_served) dt per load group with policy-independent demand"],
       ["LOLH", "Σ 1[P_unserved > ε_P]·Δt, ε_P = 0.5 W"],
       ["LPSP_E", "E_unserved / E_demand (energy-based)"],
       ["Availability", "successful / valid scheduled HTTP probes (1 s grid)"],
       ["RTT", "command sent → ACK received (PC clock)"],
       ["Decision → effect", "EMS decision → first 10 Hz M4 sample showing the new state"],
       ["Time to limit", "RUN_START → first SOC_ref ≤ 0.20; right-censored at 6 h"]],
      widths=[3.6, 13.4], size=8.5)
P("Statistical design: the unit is a paired block (same bench unit, initial state and noise seed for all policies). Paired differences are reported "
  "with a 95% BCa bootstrap interval (10,000 resamples) and an exact two-sided sign-flip test (2¹⁰ sign patterns, minimum p = 0.002); Holm "
  "correction over the two primary contrasts (C1−C0, C2−C1); Friedman omnibus across C0/C1/C2. Secondary outcomes are exploratory.")

# ============================================================ PART D
page_break()
H("Part D. Results")
H("D.1. Acceptance summary", 2)
table("Protocol acceptance criteria (§17) on the virtual testbed",
      ["Object", "Criterion", "Observed", "Verdict"],
      [[a["object"], a["criterion"], a["observed"], "met" if a["pass"] else ("see D.7" if a["pass"] is None else "NOT met")] for a in S["acceptance"]],
      widths=[3.2, 4.4, 7.2, 2.2], size=8)
table("Key results at a glance",
      ["Quantity", "Value"],
      [["Runs / blocks", f"{MAN['n_runs']} runs; E05 {E5['blocks']} paired blocks × 4 policies"],
       ["SOC accuracy (nominal runs)", f"MAE ≤ {f(MAX_MAE)} pp; max |error| ≤ {f(max(r['max_soc_err_pp'] for r in NR))} pp"],
       ["Energy error", f"≤ {f(MAX_EN)} %"],
       ["Command RTT", f"median {f(T['rtt_ms']['median'], 1)} ms; p95 {f(T['rtt_ms']['p95'], 1)} ms; max {f(T['rtt_ms']['max'], 1)} ms (n = {T['rtt_ms']['n']})"],
       ["Decision → current edge", f"median {f(T['effect_ms']['median'], 0)} ms; max {f(T['effect_ms']['max'], 0)} ms (n = {T['effect_ms']['n']})"],
       ["Confirmed executions", f"{sum(r['commands_effect_confirmed'] for r in S['runs'])}/{sum(r['commands_accepted'] for r in S['runs'])}"],
       ["Fault detection", f"{sum(r['detected'] for r in S['e07']['table'])}/5, false flags {S['e07']['false_flags']}"],
       ["E05: C1 − C0 time to limit", f"+{f(E5['primary']['C1_minus_C0']['mean'])} h (Holm p = {f(E5['primary']['C1_minus_C0']['p_holm'], 4)})"],
       ["E05: C2 − C1 lighting", f"{f(E5['secondary']['aux_served_wh:C2-C1']['mean'], 1)} Wh (H4 not supported)"],
       ["E11: critical unserved, S1 100 Ah", f"C0 {f(E11['controller_replay']['C0']['e_unserved_critical_wh'], 0)} → C2 {f(E11['controller_replay']['C2']['e_unserved_critical_wh'], 0)} Wh"]],
      widths=[5.5, 11.5], size=8.5)

H("D.2. Metrology", 2)
figure("fig09_calibration", "Residuals after E00 calibration: circles — fit levels, diamonds — verification levels. The ±1 % target is off scale.", 15.5)
table("E00 calibration results (EMU)",
      ["Channel", "Raw zero, mA", "a", "b, mA", "Max verification residual, %", "Pass"],
      [[n, f(c["zero_offset_raw_a"] * 1e3, 2), f(c["a_i"], 5), f(c["b_i"] * 1e3, 2), f(c["max_verification_rel"] * 100, 3), "yes" if c["pass_current"] else "NO"]
       for n, c in S["calibration"].items()], widths=[2.0, 2.6, 2.4, 2.2, 5.0, 2.8])
figure("fig12_uncertainty_budget", "Uncertainty budget of calibrated 1-s current values: variance shares of the components (left) and expanded uncertainty U_I (k = 2) (right).")
table("Expanded uncertainty (k = 2) at operating points",
      ["Channel @ I", "U_I, mA", "U_I, % of reading", "U_V, %", "U_P, %"],
      [[f"{ch} @ {f(r['i_op_a'], 0)} A", f(r["U_i_a_k2"] * 1e3, 2), f(abs(r["U_i_rel_k2"]) * 100, 3), f(r["U_v_rel_k2"] * 100, 3), f(abs(r["U_p_rel_k2"]) * 100, 3)]
       for ch, rows in S["uncertainty"].items() for r in rows], widths=[3.6, 2.6, 3.6, 3.4, 3.8])

H("D.3. Scenario dynamics", 2)
figure("fig11_scenarios", "Dynamics of E02 (step load), E03 (generation drop, C1), E04 (source loss) and E08 (recovery, C1): bus powers (top), "
       "battery current (middle, + = discharge) and SOC_ref vs twin posterior (bottom).", 17.0)
table("Per-scenario outcomes (single runs; E03/E08 per policy)",
      ["Run", "Duration, h", "Stop", "Coverage, %", "Min SOC_ref", "Lighting served, Wh", "Switches", "Commands (confirmed)", "Availability, %"],
      [[f"{r['scenario_id']} {r['policy']}", f(r["duration_h"]), r["stop_reason"], f(r["coverage_pct"]), f(r["min_soc_ref"], 3), f(r["aux_served_wh"], 1),
        r["switch_count"], f"{r['commands']} ({r['commands_effect_confirmed']})", f(r["availability_pct"], 3)] for r in S["runs"] if r["scenario_id"] != "E05"],
      widths=[1.8, 1.6, 2.2, 1.8, 1.8, 2.2, 1.4, 2.2, 2.0], size=7.8)

H("D.4. Twin accuracy (H1)", 2)
figure("fig04_twin_accuracy", "SOC tracking (A) and error (B) for E03-C2 and E04-C0; shaded band = ±3 pp criterion.")
table("State and energy accuracy",
      ["Run", "MAE, pp", "RMSE, pp", "Max, pp", "Energy error, %", "REF − truth, pp", "r_P mean ± SD, W"],
      [[f"{r['scenario_id']} {r['policy']}", f(r["mae_soc_pp"]), f(r["rmse_soc_pp"]), f(r["max_soc_err_pp"]),
        f(r["energy_error_pct"]) if r["energy_error_pct"] == r["energy_error_pct"] else "—", f(r["ref_vs_true_mae_pp"]),
        f"{f(r['residual_w_mean'])} ± {f(r['residual_w_sd'])}"] for r in NR if r["scenario_id"] != "E05"] +
      [[f"E05 {p} (n=10, mean)", f(PP[p]["mae_soc_pp"]["mean"]), "—", f(PP[p]["max_soc_err_pp"]["max"]), f(PP[p]["energy_error_pct"]["mean"]), "—", "—"]
       for p in ("C0", "C1", "C2", "C_oracle")], widths=[3.2, 1.8, 1.8, 1.8, 2.6, 2.6, 3.2], size=8)

H("D.5. Synchronisation and execution (H2, H3)", 2)
figure("fig08_timing", "Command → ACK RTT (A) and decision → M4 current edge (B).", 16.0)

H("D.6. Faults and protection logic", 2)
figure("fig05_e06_link_loss", "E06: ESP32 output and measured relay state during uplink loss (600–630 s) and bidirectional loss (900–930 s); stale and foreign-run commands rejected at 1200 s.", 16.0)
e6 = S["e06"]
table("E06 event log",
      ["Time, s", "Event / controller entry"],
      [[f(e["t"], 1), e["event_type"]] for e in e6["events"]] + [[f(r[0], 1), f"{r[1]} {r[2]}"] for r in e6["edge"]],
      widths=[2.5, 14.5], size=8)
FU = {"offset": "current offset", "stuck": "stuck value", "missing": "missing field", "impossible": "impossible value", "sign": "sign flip"}
table("E07 fault detection",
      ["Fault", "Parameter", "Duration, s", "Detected", "Latency, s", "Flag", "Flagged share"],
      [[FU[r["fault"]], r["param"] if r["param"] is not None else "—", r["duration_s"], "yes" if r["detected"] else "NO",
        f(r["latency_s"], 2) if r["latency_s"] is not None else "—", r["flag"], f(100 * r["flagged_fraction"], 0) + " %"] for r in S["e07"]["table"]],
      widths=[2.4, 2.4, 1.8, 1.6, 1.8, 4.0, 2.2], size=8)
table("E10 protection and contract vectors",
      ["Vector", "Expected", "Observed", "State after", "Result"],
      [[r["test"], "yes" if r["expected_accepted_or_condition"] else "no", "yes" if r["observed"] else "no", r["state_after"], "PASS" if r["pass"] else "FAIL"]
       for r in S["e10"]], widths=[8.0, 1.8, 2.0, 2.6, 2.0], size=8)

H("D.7. Policy comparison in the deficit test (E05, H4)", 2)
figure("fig02_e05_timeseries", "E05 block R00: SOC_ref of the four policies (A); available generation and lighting state (B).")
figure("fig03_e05_paired", "E05 paired results over 10 blocks: time to limit (A), lighting delivered (B), minimum SOC_ref (C).")
table("E05 per-policy summary (mean ± SD over 10 blocks)",
      ["Policy", "Time to limit, h", "Censored", "Lighting, Wh", "Min SOC_ref", "Switches", "SOC MAE, pp"],
      [[p, f"{f(PP[p]['time_to_limit_h']['mean'])} ± {f(PP[p]['time_to_limit_h']['sd'])}", f"{PP[p]['censored']}/10",
        f"{f(PP[p]['aux_served_wh']['mean'], 1)} ± {f(PP[p]['aux_served_wh']['sd'], 1)}", f"{f(PP[p]['min_soc_ref']['mean'], 3)} ± {f(PP[p]['min_soc_ref']['sd'], 3)}",
        f(PP[p]["switch_count"]["mean"], 1), f(PP[p]["mae_soc_pp"]["mean"])] for p in ("C0", "C1", "C2", "C_oracle")],
      widths=[1.8, 3.0, 1.6, 3.0, 3.0, 1.6, 2.0], size=8.5)
rows = []
for k, lab in (("C1_minus_C0", "Time to limit C1 − C0, h"), ("C2_minus_C1", "Time to limit C2 − C1, h")):
    v = E5["primary"][k]
    rows.append([lab + " (primary)", f(v["mean"], 3), f"{f(v['ci_lo'], 3)}; {f(v['ci_hi'], 3)}", f(v["p_signflip"], 4), f(v["p_holm"], 4)])
for k, v in E5["secondary"].items():
    rows.append([k.replace(":", ": ").replace("_", " "), f(v["mean"], 3), f"{f(v['ci_lo'], 3)}; {f(v['ci_hi'], 3)}", f(v["p_signflip"], 4), "expl."])
table("E05 paired contrasts (improvement orientation: later policy − earlier policy)",
      ["Contrast", "Mean", "95% BCa CI", "p sign-flip", "p Holm"], rows, widths=[7.4, 2.0, 3.6, 2.0, 2.0], size=7.8,
      note=f"Friedman across C0/C1/C2: lighting χ² = {f(E5['friedman']['aux_served_wh']['chi2'], 1)}, p = {E5['friedman']['aux_served_wh']['p']:.1e}. "
           "Intervals reflect measurement-noise variability on one bench unit only.")

H("D.8. Fourteen-day scenario and reproduction of the manuscript (E11)", 2)
figure("fig06_e11_14day", "E11: stored-energy fraction for S1 and S3 (A); critical and lighting unserved energy for S1 with 12.8 V / 100 Ah under C0–C2 (B).")
art = {"S1": 0.219, "S2": 0.219, "S3": 0.247}
sil = E11["sil"]
table("E11 SIL reproduction",
      ["Scenario", "MinSOC manuscript", "SIL 1 h", "SIL 1 min", "Nameplate V", "LOLH, h", "LPSP_E", "Curtailed, Wh"],
      [[s, f(art[s], 3), f(sil[s]["min_soc"], 3), f(E11["sil_1min"][s]["min_soc"], 3), f(sil[s]["min_soc_nameplate_v"], 3), f(sil[s]["lolh_h"], 0),
        f(sil[s]["lpsp_e"], 4), f(sil[s]["e_curtailed_wh"], 0)] for s in ("S1", "S2", "S3")], widths=[1.8, 2.4, 1.8, 1.8, 2.2, 1.6, 1.8, 2.4], size=8.5)
cr = E11["controller_replay"]
table("E11 lockstep controller replay, S1 with 12.8 V / 100 Ah",
      ["Policy", "Critical unserved, Wh", "Lighting unserved, Wh", "Total bus unserved, Wh", "LOLH, h", "Commands accepted", "Switches"],
      [[p, f(cr[p]["e_unserved_critical_wh"], 0), f(cr[p]["e_unserved_aux_wh"], 0), f(cr[p]["e_unserved_bus_wh"], 0), f(cr[p]["lolh_h"], 0),
        f"{cr[p]['commands']['accepted']}/{cr[p]['commands']['sent']}", cr[p]["switches"]] for p in ("C0", "C1", "C2")],
      widths=[1.6, 2.8, 2.8, 2.8, 1.8, 2.8, 2.4], size=8.5)
table("MinSOC sensitivity to undocumented manuscript parameters",
      ["Day order", "Battery efficiency", "k_T", "S1", "S3"],
      [[r["order"], r["efficiency"], f(r["k_t"]), f(r["min_soc_S1"], 3), f(r["min_soc_S3"], 3)] for r in E11["reconciliation"]],
      widths=[4.0, 6.4, 1.4, 2.6, 2.6], size=8)

H("D.9. Analytical sizing and sensitivity", 2)
table("Sizing (B.4–B.7) — exact agreement with the manuscript",
      ["Scenario", "E_load, Wh/d", "E_src, Wh/d", "E_batt, Wh", "C, Ah", "P_PV, Wp", "I²R vs 12 V", "Mean bus I, A"],
      [[s, f(v["e_load_day"], 0), f(v["e_src_day"], 1), f(v["e_batt_req"], 1), f(v["c_ah"], 1), f(v["p_pv"], 1), f(v["loss_ratio_vs_12v"], 4), f(v["i_bus_mean_a"], 2)]
       for s, v in S["sizing"].items()], widths=[1.6, 2.2, 2.2, 2.2, 1.6, 2.0, 2.4, 2.8], size=8.5)
table("Sensitivity (S1): required Wp and Ah (reproduces Table 4 of the manuscript)",
      ["η_dc", "PSH 1.2, Wp", "PSH 1.8, Wp", "PSH 2.5, Wp", "C, Ah"],
      [[f(r["eta_dc"]), f(r["wp_psh_1.2"], 0), f(r["wp_psh_1.8"], 0), f(r["wp_psh_2.5"], 0), f(r["c_ah"], 1)] for r in S["sensitivity"]],
      widths=[2.4, 3.4, 3.4, 3.4, 3.4])
figure("fig07_sensitivity", "Required PV power for S1 versus PSH and η_dc.", 11.0)

# ============================================================ PART E
page_break()
H("Part E. Manuscript integration kit")
H("E.1. Proposed new subsection (ready to paste)", 2)
P("Suggested title: **“Validation of the digital twin on a virtual testbed with emulated hardware”** (keep “hardware-in-the-loop” out of the title "
  "until physical runs exist).", italic=True)
e5c1 = E5["primary"]["C1_minus_C0"]
paras = [
    ("Testbed.", f"The digital twin was validated on a virtual testbed that reproduces a 12 V communication-node power supply built from commercially "
     f"available components: a 12.8 V / 40 Ah LiFePO₄ pack with BMS, a programmable DC supply (RIDEN RD6018) acting as PV equivalent, Mean Well "
     f"DDR-60G-12 converters feeding an 18 W critical load (router and electronic load) and a relay-switched 30 W lighting branch, INA226 monitors on "
     f"FL-2 shunts at six measurement points and an independent INA228 reference channel, and an ESP32-S3 edge controller connected over MQTT "
     f"(Fig. X, Table X). Each component was emulated with its datasheet characteristics (ranges, LSBs, offset and gain errors, noise, delays), "
     f"while the twin and the energy-management system run unchanged software."),
    ("Protocol.", f"Channels were calibrated against a reference standard (E00) and the working-interval capacity was measured by the reference "
     f"channel (E09). Scenarios covered steady state, step disturbances, source loss, a controlled energy deficit, link loss, faulty sensors, recovery, "
     f"protection logic and the 14-day stress profile of the original study. Three policies — schedule (C0), SOC threshold (C1) and a forecast-based "
     f"rule (C2) — were compared in {E5['blocks']} paired blocks with an information-bound oracle."),
    ("Results.", f"The twin tracked SOC with a mean absolute error of at most {f(MAX_MAE)} pp against the independent reference and reproduced the "
     f"discharge energy within {f(MAX_EN)} %. Command round-trip time was {f(T['rtt_ms']['p95'], 0)} ms at the 95th percentile, and every accepted "
     f"command produced a measured current change (median {f(T['effect_ms']['median'], 0)} ms after the decision). Link loss and all five injected "
     f"data faults were detected and handled without interrupting the critical load. Threshold control extended operation to the working limit by "
     f"{f(e5c1['mean'])} h (95% CI {f(e5c1['ci_lo'])}–{f(e5c1['ci_hi'])}), whereas the forecast rule with a persistence forecast did not improve "
     f"the primary outcome and delivered {f(-E5['secondary']['aux_served_wh:C2-C1']['mean'], 1)} Wh less lighting; the oracle bound indicates that "
     f"the benefit of predictive control is limited by the generation forecast."),
    ("Limitations.", "These results were obtained with emulated hardware and establish software and protocol readiness, not physical performance. "
     "A physical campaign on the specified bench, using the same software and measurement protocols, is required before hardware-in-the-loop or "
     "field claims can be made."),
]
for head, txt in paras:
    P(f"**{head}** {txt}")
H("E.2. Figure and table mapping", 2)
table("Mapping of dossier figures to the manuscript",
      ["Dossier figure", "File", "Suggested manuscript placement", "Short caption"],
      [["Architecture", "fig01_architecture.pdf", "Methods, new subsection — Fig. 13", "Architecture of the validation testbed"],
       ["Wiring", "fig10_bench_wiring.pdf", "Methods — Fig. 14 or Supplementary", "Wiring and measurement points of the 12 V bench"],
       ["Calibration", "fig09_calibration.pdf", "Supplementary", "Calibration residuals of the measurement channels"],
       ["Uncertainty", "fig12_uncertainty_budget.pdf", "Supplementary", "Uncertainty budget of current measurements"],
       ["Scenario dynamics", "fig11_scenarios.pdf", "Results — Fig. 15", "Response of the twin to load, generation and source disturbances"],
       ["Twin accuracy", "fig04_twin_accuracy.pdf", "Results — Fig. 16", "SOC tracking against the independent reference"],
       ["Timing", "fig08_timing.pdf", "Results or Supplementary", "Command and actuation latency"],
       ["Link loss", "fig05_e06_link_loss.pdf", "Supplementary", "Behaviour under communication loss"],
       ["Policy comparison", "fig02_e05_timeseries.pdf + fig03_e05_paired.pdf", "Results — Fig. 17 (two panels)", "Comparison of control policies under deficit"],
       ["14-day", "fig06_e11_14day.pdf", "Results — Fig. 18", "14-day stress profile and load-priority redistribution"],
       ["Pipeline", "fig13_data_pipeline.pdf", "Data availability / Supplementary", "Reproducible data pipeline"]],
      widths=[2.8, 4.6, 4.6, 5.0], size=8)
H("E.3. Insertions for other manuscript sections", 2)
table("Proposed text insertions",
      ["Section", "Insertion"],
      [["Highlights", f"“A virtual testbed with emulated real components validates the twin: SOC error ≤ {f(MAX_MAE)} pp, actuation confirmed for every command.”"],
       ["Abstract (append)", "“The twin was further validated on a virtual testbed emulating a 12 V bench of real components, showing accurate state tracking, reliable closed-loop actuation and robust fault handling; threshold-based load control extended autonomy, while forecast-based control was limited by forecast quality.”"],
       ["Methods", "Insert Part B text (B.1–B.8) as subsection “Experimental testbed” and Part C as “Validation protocol and metrics”."],
       ["Results", "Insert E.1 “Results” paragraph with Tables D.1/D.7 and Figs. 15–18."],
       ["Discussion", "State that predictive control requires a forecast better than persistence and that bench conclusions do not scale automatically to S1–S3."],
       ["Conclusions", "Add: software and protocols are ready for physical validation; physical HIL campaign is the next step."],
       ["Data availability", "“Code, configurations, raw logs of all 51 runs, protocols and figures are available in the dt_power_hil package (repository and DOI to be added). No physical measurements are included.”"]],
      widths=[3.2, 13.8], size=8.5)
H("E.4. Response to the reviewer (draft)", 2)
P("“We thank the reviewer for requesting hardware-based validation. We have specified a 12 V testbed on commercially available components and "
  "implemented the complete software chain (digital twin, energy management, ESP32 edge controller logic, logging and analysis). Before spending "
  "battery life on physical runs, we validated this chain on a virtual testbed in which each component is emulated with its datasheet "
  "characteristics. The revised manuscript reports these results explicitly as a virtual-testbed validation (new Section X, Figs. X–Y). "
  "Physical hardware-in-the-loop measurements will be reported once the campaign described in the measurement protocols has been completed.”", italic=True)
H("E.5. Claims that are and are not supported", 2)
table("Claim register",
      ["Supported by this dossier", "Not supported — do not claim"],
      [["Software chain works end-to-end with realistic instrument errors", "Hardware-in-the-loop, CHIL, PHIL or field validation"],
       ["SOC error ≤ 3 pp against an independent reference in emulation", "Physical SOC accuracy of the bench"],
       ["Every accepted command had a measured effect (emulated)", "Real Wi-Fi/MQTT latency of the bench"],
       ["Threshold control extends autonomy on the 40 Ah bench model", "Superiority of predictive (C2/MPC/RL) control"],
       ["Analytical sizing of S1–S3 reproduced exactly", "Exact reproduction of MinSOC 0.219 / 0.247"],
       ["Protocols and tolerances are derived from component datasheets", "Battery degradation, seasonal reliability or 48 V field performance"]],
      widths=[8.5, 8.5], size=8.5)

# ============================================================ APPENDIX
page_break()
H("Appendix. Full per-run metrics")
table("All 51 runs (EMU)",
      ["run_id", "h", "Stop", "Cov. %", "MAE", "Max", "E err %", "Light Wh", "Min SOC_ref", "RTT p95 ms"],
      [[r["run_id"].replace("20260929_", "").replace("_EMU", ""), f(r["duration_h"]), "limit" if r["stop_reason"] == "stopped_limit" else "done",
        f(r["coverage_pct"]), f(r["mae_soc_pp"]), f(r["max_soc_err_pp"]), f(r["energy_error_pct"]) if r["energy_error_pct"] == r["energy_error_pct"] else "—",
        f(r["aux_served_wh"], 1), f(r["min_soc_ref"], 3), f(r["rtt"].get("p95_ms", float("nan")), 1) if r["rtt"]["n"] else "—"] for r in S["runs"]],
      widths=[3.4, 1.2, 1.3, 1.5, 1.2, 1.2, 1.5, 1.6, 2.0, 2.0], size=7.5, mono_cols=(0,))
cal = json.load(open(RES / "calibration" / "E00_calibration.json"))
for n in ("M1", "REF"):
    table(f"E00 calibration levels, channel {n}",
          ["Level, A", "Standard, A", "DUT raw, A", "DUT SD, mA", "Calibrated, A", "Residual, mA", "Residual, %", "Role"],
          [[f(r["level_a"], 2), f(r["i_std"], 4), f(r["i_dut_mean"], 4), f(r["i_dut_sd"] * 1e3, 2), f(r["i_cal"], 4), f(r["residual_a"] * 1e3, 2),
            f(r["residual_rel"] * 100, 3), "fit" if k in (0, 2, 4) else "verify"] for k, r in enumerate(cal[n]["levels"])],
          widths=[1.8, 2.2, 2.2, 1.9, 2.2, 2.2, 2.0, 1.5], size=8)

doc.core_properties.title = "Experiment Dossier — Digital Twin Validation Testbed"
doc.core_properties.subject = S["evidence_class"]
for z in doc.settings.element.iter(qn("w:zoom")):
    z.set(qn("w:percent"), "100")
out = ROOT / "report" / "DT_HIL_Experiment_Dossier_Testbed_EN.docx"
doc.save(out)
print("saved", out, f"{os.path.getsize(out) / 1e6:.1f} MB", FIGN[0], "figures", TABN[0], "tables")
